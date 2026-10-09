from __future__ import annotations

import re
from pathlib import Path
from typing import Any

import snowballstemmer
from langchain_core.callbacks import CallbackManagerForRetrieverRun
from langchain_core.documents import Document
from langchain_core.embeddings import Embeddings
from langchain_core.retrievers import BaseRetriever
from langchain_core.vectorstores import InMemoryVectorStore
from langchain_text_splitters import RecursiveCharacterTextSplitter
from rank_bm25 import BM25Okapi

DOC_EXTENSIONS = {".txt", ".md", ".pdf", ".docx"}

_stemmers = {"ru": snowballstemmer.stemmer("russian"), "en": snowballstemmer.stemmer("english")}
_word_re = re.compile(r"\w+", re.UNICODE)
_cyrillic_re = re.compile("[а-яё]")
_lab_number_re = re.compile(r"(?:\bлр|работ\w*)\s*[-№]?\s*(?:номер\s*)?(\d+)", re.IGNORECASE)


def tokenize(text: str) -> list[str]:
    words = _word_re.findall(text.lower())
    return [_stemmers["ru" if _cyrillic_re.search(w) else "en"].stemWord(w) for w in words]


def load_file(path: str | Path) -> list[Document]:
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader

        pages = [(page.extract_text() or "").strip() for page in PdfReader(str(path)).pages]
        return [
            Document(page_content=text, metadata={"source": path.name, "page": number})
            for number, text in enumerate(pages, start=1)
            if text
        ]
    if ext == ".docx":
        import docx

        document = docx.Document(str(path))
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            parts += [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        text = "\n".join(parts)
    elif ext in {".txt", ".md"}:
        text = path.read_text(encoding="utf-8", errors="ignore")
    else:
        raise ValueError(f"Формат {ext} не поддерживается")
    return [Document(page_content=text, metadata={"source": path.name})] if text.strip() else []


def format_docs(docs: list[Document]) -> str:
    return "\n\n".join(f"[{d.metadata.get('source', '?')}]\n{d.page_content}" for d in docs)


class HybridRetriever(BaseRetriever):
    kb: Any
    k: int = 4

    def _get_relevant_documents(
        self, query: str, *, run_manager: CallbackManagerForRetrieverRun
    ) -> list[Document]:
        return self.kb.search(query, self.k)


class KnowledgeBase:
    def __init__(
        self, embeddings: Embeddings | None = None, chunk_size: int = 900, chunk_overlap: int = 150
    ) -> None:
        self.embeddings = embeddings
        self.splitter = RecursiveCharacterTextSplitter(
            chunk_size=chunk_size,
            chunk_overlap=chunk_overlap,
            separators=["\n## ", "\n### ", "\n\n", "\n", ". ", " ", ""],
        )
        self.chunks: list[Document] = []
        self.titles: dict[str, str] = {}
        self.vector_store: InMemoryVectorStore | None = None
        self.vector_error = ""
        self._bm25: BM25Okapi | None = None

    def add_directory(self, directory: Path) -> str:
        files = sorted(p for p in Path(directory).glob("*") if p.suffix.lower() in DOC_EXTENSIONS)
        return self.add_files(files) if files else "Каталог методичек пуст"

    def add_files(self, paths: list[str | Path]) -> str:
        report, new_chunks = [], []
        for path in map(Path, paths):
            try:
                docs = load_file(path)
            except Exception as exc:
                report.append(f"{path.name}: ошибка чтения ({exc})")
                continue
            if not docs:
                report.append(f"{path.name}: текст не найден")
                continue
            self.chunks = [c for c in self.chunks if c.metadata["source"] != path.name]
            title = self._title(docs[0].page_content, path)
            self.titles[path.name] = title
            chunks = self.splitter.split_documents(docs)
            for number, chunk in enumerate(chunks, start=1):
                chunk.metadata.update(title=title, chunk=number)
                if title not in chunk.page_content[: len(title) + 4]:
                    chunk.page_content = f"{title}\n{chunk.page_content}"
            new_chunks += chunks
            report.append(f"{path.name}: {len(chunks)} фрагм.")
        self.chunks += new_chunks
        self._rebuild_index()
        return "; ".join(report)

    def _rebuild_index(self) -> None:
        self._bm25 = (
            BM25Okapi([tokenize(c.page_content) for c in self.chunks]) if self.chunks else None
        )
        self.vector_store = None
        if self.embeddings is None or not self.chunks:
            return
        try:
            store = InMemoryVectorStore(self.embeddings)
            store.add_documents(self.chunks, ids=[str(i) for i in range(len(self.chunks))])
            self.vector_store, self.vector_error = store, ""
        except Exception as exc:
            self.vector_error = str(exc)

    @staticmethod
    def _title(text: str, path: Path) -> str:
        for line in text.splitlines():
            if line.strip():
                return line.strip("# ").strip()[:120]
        return path.stem

    def search(self, query: str, k: int = 4) -> list[Document]:
        if not self.chunks:
            return []
        rankings: list[list[int]] = []
        depth = max(3 * k, 10)
        scores = self._bm25.get_scores(tokenize(query))
        bm25_rank = [
            i for i in sorted(range(len(scores)), key=lambda i: -scores[i]) if scores[i] > 0
        ]
        rankings.append(bm25_rank[:depth])
        numbers = set(_lab_number_re.findall(query))
        if numbers:
            wanted = {
                i
                for i, c in enumerate(self.chunks)
                if any(f"№ {n}." in c.metadata.get("title", "") for n in numbers)
            }
            rankings.append([i for i in bm25_rank if i in wanted] + sorted(wanted - set(bm25_rank)))
        if self.vector_store is not None:
            try:
                found = self.vector_store.similarity_search(query, k=depth)
                rankings.append([int(doc.id) for doc in found])
            except Exception as exc:
                self.vector_error = str(exc)
        fused: dict[int, float] = {}
        for ranking in rankings:
            for rank, i in enumerate(ranking):
                fused[i] = fused.get(i, 0.0) + 1.0 / (60 + rank + 1)
        best = sorted(fused, key=fused.get, reverse=True)[:k]
        return [self.chunks[i] for i in best]

    def as_retriever(self, k: int = 4) -> HybridRetriever:
        return HybridRetriever(kb=self, k=k)

    def describe(self) -> str:
        if not self.chunks:
            return "База знаний пуста: загрузите методические указания."
        mode = (
            "векторный + BM25" if self.vector_store is not None else "BM25 (эмбеддинги недоступны)"
        )
        files = "\n".join(f"- {title}" for title in self.titles.values())
        stats = f"Документов: {len(self.titles)}, фрагментов: {len(self.chunks)}, поиск: {mode}"
        return f"{stats}\n\n{files}"
