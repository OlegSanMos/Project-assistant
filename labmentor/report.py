from __future__ import annotations

import re
import uuid
from pathlib import Path

import pandas as pd
from docx import Document
from docx.shared import Cm, Pt

from .formulas import TAGS, latex_to_text


def _add_text(paragraph, text: str, bold: bool = False) -> None:
    for part in re.split(r"(\*\*[^*]+\*\*|<su[bp]>.*?</su[bp]>)", text):
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            _add_text(paragraph, part[2:-2], bold=True)
        elif part:
            run = paragraph.add_run(TAGS.sub("", part))
            run.bold = bold or None
            if part.startswith("<sub>"):
                run.font.subscript = True
            elif part.startswith("<sup>"):
                run.font.superscript = True


def _add_measurements(document, table: pd.DataFrame | None, plot: str | None) -> None:
    if table is not None and len(table):
        labels = table.attrs.get("labels", {})
        document.add_paragraph("Таблица 1. Результаты измерений")
        grid = document.add_table(rows=1, cols=len(table.columns))
        grid.style = "Table Grid"
        for cell, column in zip(grid.rows[0].cells, table.columns, strict=True):
            cell.text = labels.get(column, column)
        for _, row in table.head(50).iterrows():
            for cell, value in zip(grid.add_row().cells, row, strict=True):
                cell.text = f"{value:.6g}" if isinstance(value, float) else str(value)
    if plot and Path(plot).exists():
        document.add_picture(plot, width=Cm(15))
        document.add_paragraph("Рисунок 1. Экспериментальная зависимость и аппроксимация МНК")


def export_report_docx(
    markdown: str,
    title: str,
    out_dir: Path,
    table: pd.DataFrame | None = None,
    plot: str | None = None,
) -> str:
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Times New Roman", Pt(14)
    document.add_heading("Отчет по лабораторной работе", level=0)
    document.add_paragraph().add_run(title).bold = True
    inserted = False
    for line in latex_to_text(markdown, markup=True).splitlines():
        line = line.rstrip()
        if not line.strip() or line.startswith("# "):
            continue
        heading = re.match(r"^(#{2,4})\s+(.*)", line)
        bullet = re.match(r"^\s*[-*]\s+(.*)", line)
        number = re.match(r"^\s*\d+[.)]\s+(.*)", line)
        if heading:
            title_text = TAGS.sub("", heading.group(2)).strip()
            document.add_heading(title_text, level=len(heading.group(1)) - 1)
            if not inserted and "результат" in title_text.lower():
                _add_measurements(document, table, plot)
                inserted = True
        elif bullet:
            _add_text(document.add_paragraph(style="List Bullet"), bullet.group(1))
        elif number:
            _add_text(document.add_paragraph(style="List Number"), number.group(1))
        else:
            _add_text(document.add_paragraph(), line)
    if not inserted:
        _add_measurements(document, table, plot)
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"report_{uuid.uuid4().hex[:8]}.docx"
    document.save(str(path))
    return str(path)
