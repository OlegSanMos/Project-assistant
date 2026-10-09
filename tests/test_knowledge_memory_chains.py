from pathlib import Path

import docx
from langchain_core.embeddings import DeterministicFakeEmbedding
from langchain_core.runnables import RunnableLambda

from labmentor import chains
from labmentor.knowledge import KnowledgeBase, load_file, tokenize
from labmentor.memory import ConversationMemory

MANUALS = Path(__file__).resolve().parent.parent / "data" / "manuals"


def test_tokenize_uses_russian_stemming():
    assert tokenize("Сопротивления сопротивлением") == ["сопротивлен", "сопротивлен"]


def test_bm25_search_finds_relevant_manual():
    kb = KnowledgeBase(None)
    kb.add_directory(MANUALS)
    assert len(kb.titles) == 4 and kb.vector_store is None
    assert kb.search("модуль Юнга стали", 2)[0].metadata["source"] == "lr2_young_modulus.md"
    assert (
        kb.search("постоянная времени разряда конденсатора", 2)[0].metadata["source"]
        == "lr3_rc_circuit.md"
    )
    assert {d.metadata["source"] for d in kb.search("цель лабораторной работы № 1", 3)} == {
        "lr1_ohm_law.md"
    }


def test_hybrid_search_with_embeddings_and_retriever():
    kb = KnowledgeBase(DeterministicFakeEmbedding(size=32))
    kb.add_directory(MANUALS)
    assert kb.vector_store is not None
    docs = kb.as_retriever(3).invoke("закон Ома для участка цепи")
    assert len(docs) == 3 and docs[0].metadata["source"] == "lr1_ohm_law.md"


def test_docx_loading(tmp_path):
    path = tmp_path / "manual.docx"
    document = docx.Document()
    document.add_paragraph("Лабораторная работа № 7. Измерение мощности")
    document.add_paragraph("Ваттметр включают по схеме с общей точкой.")
    document.save(path)
    assert "Ваттметр" in load_file(path)[0].page_content
    kb = KnowledgeBase(None)
    assert "1 фрагм" in kb.add_files([path])


def test_context_window_trimming_and_summary():
    memory = ConversationMemory(max_tokens=120)
    for i in range(6):
        memory.add_turn(
            f"Вопрос номер {i} о законе Ома и измерениях", f"Ответ номер {i}: " + "текст " * 20
        )
    window = memory.window()
    assert 0 < len(window) < len(memory.messages) and window[0].type == "human"
    assert memory.stats()["tokens"] <= 120
    calls = []
    summary_chain = RunnableLambda(lambda x: calls.append(x) or "Обсуждали закон Ома.")
    assert memory.update_summary(summary_chain) is True
    assert (
        memory.summary_text() == "Обсуждали закон Ома." and "Вопрос номер 0" in calls[0]["dialog"]
    )
    assert memory.update_summary(summary_chain) is False
    memory.clear()
    assert memory.messages == [] and memory.summary_text() == "нет"


def test_parsers():
    assert chains.parse_route("Ответ: data") == "data"
    assert chains.parse_route("не знаю") == "manual"
    quiz = chains.parse_quiz(
        '```json\n{"questions": [{"question": "Что такое R?", "answer": "Сопротивление"}]}\n```'
    )
    assert quiz == [{"question": "Что такое R?", "answer": "Сопротивление"}]
    assert chains.parse_quiz("1. Что такое ВАХ?\n2. Как найти R?")[1]["question"] == "Как найти R?"
    assert chains.parse_grade('{"score": 14, "feedback": "ok"}')["score"] == 10
    assert chains.parse_flowchart('{"title": "Т", "steps": ["a", "b"]}')["steps"] == ["a", "b"]


def test_router_chain(scripted_llm):
    router = chains.build_router_chain(scripted_llm)
    assert (
        router.invoke({"question": "Построй график и найди погрешность", "has_table": "да"})
        == "data"
    )
    assert (
        router.invoke({"question": "Найди в интернете плотность меди", "has_table": "нет"})
        == "research"
    )
