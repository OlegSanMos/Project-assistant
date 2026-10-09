from docx import Document

from conftest import FORMULA_ANSWER
from labmentor import chains, speech
from labmentor.assistant import Reply
from labmentor.formulas import latex_to_text, prettify_math
from labmentor.report import export_report_docx
from labmentor.ui import render


def test_latex_to_text_common_formulas():
    cases = {
        r"$R = \frac{U}{I}$": "R = U/I",
        r"$$\sigma = \frac{F}{S}$$": "\nσ = F/S\n",
        r"\(\tau = R \cdot C = 10^{3} \cdot 10^{-3}\ \text{с}\)": "τ = R · C = 10³ · 10⁻³ с",
        r"\[\Delta R = t_{0,95} \frac{S}{\sqrt{n}}\]": "\nΔR = t_0,95 S/√n\n",
        r"$I_1 = \frac{U_1 - U_2}{R_x}$": "I₁ = (U₁ - U₂)/Rₓ",
        r"$\bar{x} = \frac{1}{n}\sum_{i=1}^{n} x_i$": "x̄ = 1/n Σᵢ₌₁ⁿ xᵢ",
        r"$T = 20^\circ\text{C}$, $R = 101{,}5\,\text{Ом}$": "T = 20°C, R = 101,5 Ом",
        r"$u(t) = U_0 e^{-t/\tau}$, $\ln 2 \approx 0{,}693$": "u(t) = U₀ e^(-t/τ), ln 2 ≈ 0,693",
        r"$a = \sqrt[3]{V}$, $\vec{F} = m\vec{a}$": "a = ∛V, F⃗ = ma⃗",
        r"$$\begin{aligned} U &= I R \\ P &= U I \end{aligned}$$": "\nU = I R; P = U I\n",
        r"$\left. \frac{dU}{dt} \right|_{t=0} = -\frac{U_0}{\tau}$": "dU/dt |ₜ₌₀ = -U₀/τ",
    }
    for latex, text in cases.items():
        assert latex_to_text(latex) == text
    assert latex_to_text("Цена 5 $ без формул") == "Цена 5 $ без формул"


def test_latex_markup_for_docx_indices():
    text = latex_to_text(r"$R_{\text{ср}} = 101{,}5\ \text{Ом}$, $e^{-t/\tau_0}$", markup=True)
    assert text == "R<sub>ср</sub> = 101,5 Ом, e<sup>-t/τ₀</sup>"


def test_prettify_math_fixes_decimal_comma_only_in_formulas():
    text = r"Получено $R = 101,5\ \text{Ом}$, в 1,5 раза больше; $$t = 0,95$$"
    assert prettify_math(text) == (
        r"Получено $R = 101{,}5\ \text{Ом}$, в 1,5 раза больше; $$t = 0{,}95$$"
    )
    assert prettify_math(prettify_math(text)) == prettify_math(text)
    assert render(Reply(text="$x = 0,5$"))[-1]["content"] == "$x = 0{,}5$"


def test_speech_reads_formulas_as_text():
    spoken = speech.text_for_speech(FORMULA_ANSWER)
    assert "\\" not in spoken and "$" not in spoken
    assert "I = U/R" in spoken and "P = U · I ≤ 0,5 Вт" in spoken and "τ = R н C" in spoken


def test_docx_report_formulas_with_indices(tmp_path):
    markdown = (
        "## Теория\nЗакон Ома: $R = \\frac{U}{I}$.\n"
        "$$R_{\\text{ср}} = 101{,}5\\ \\text{Ом}, \\quad \\tau = 10^{-3}\\ \\text{с}$$\n"
        "**Итог: $I_{\\max} = 70$ мА**"
    )
    document = Document(export_report_docx(markdown, "ЛР 1", tmp_path))
    texts = [p.text for p in document.paragraphs]
    assert "Закон Ома: R = U/I." in texts and "Rср = 101,5 Ом, τ = 10-3 с" in texts
    runs = next(p for p in document.paragraphs if p.text.startswith("Rср")).runs
    assert [r.text for r in runs if r.font.subscript] == ["ср"]
    assert [r.text for r in runs if r.font.superscript] == ["-3"]
    runs = next(p for p in document.paragraphs if p.text.startswith("Итог")).runs
    assert "".join(r.text for r in runs) == "Итог: Imax = 70 мА" and all(r.bold for r in runs)
    assert [r.text for r in runs if r.font.subscript] == ["max"]


def test_prompts_ask_for_latex_except_json_answers():
    assert "LaTeX" in chains.ROLE
    report = chains.REPORT_PROMPT.messages[0].prompt.template
    assert "LaTeX" in report
    for prompt in (chains.ROUTER_PROMPT, chains.QUIZ_PROMPT, chains.GRADE_PROMPT):
        assert "LaTeX" not in prompt.messages[0].prompt.template
    assert "LaTeX" not in chains.FLOWCHART_PROMPT.messages[0].prompt.template
